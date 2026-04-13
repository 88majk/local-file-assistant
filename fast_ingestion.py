import os
import math
from datetime import datetime
from markitdown import MarkItDown
import ollama
from db_mongo import MongoManager
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

class FastIngestionScanner:
    def __init__(self, folder_path="data", max_workers=4):
        self.folder_path = folder_path
        self.db = MongoManager()
        self.max_workers = max_workers
        self.EMBED_MODEL = "nomic-embed-text"

    def _normalize_embedding(self, embedding):
        vec = [float(x) for x in embedding]
        norm = math.sqrt(sum(x * x for x in vec))
        return vec if norm == 0.0 else [x / norm for x in vec]

    def chunk_text(self, text, chunk_size=1000, overlap=200):
        chunks = []
        start = 0
        while start < len(text):
            chunks.append(text[start:start + chunk_size])
            start += chunk_size - overlap
        return chunks

    def _count_docx_pages(self, filepath, full_text):
        if DocxDocument is not None:
            try:
                doc = DocxDocument(filepath)
                page_breaks = 0
                for paragraph in doc.paragraphs:
                    for run in paragraph.runs:
                        page_breaks += len(run.element.xpath(".//w:br[@w:type='page']"))
                return max(1, page_breaks + 1), "pages", False
            except Exception:
                pass

        line_count = len(full_text.splitlines()) if full_text else 0
        estimated_pages = max(1, math.ceil(line_count / 45)) if line_count else 1
        return estimated_pages, "pages", True

    def _to_absolute_path(self, path):
        return os.path.normpath(os.path.abspath(path))

    def _infer_content_size(self, file_type, filepath, full_text):
        text_like_extensions = {
            "txt", "md", "markdown", "py", "js", "ts", "tsx", "jsx", "java",
            "c", "cpp", "h", "hpp", "cs", "go", "rs", "php", "rb", "swift",
            "kt", "kts", "scala", "sql", "json", "yaml", "yml", "xml", "html",
            "css", "sh", "bat", "ps1", "ini", "cfg", "toml", "csv", "log"
        }

        if file_type in {"ppt", "pptx"}:
            if Presentation is not None:
                try:
                    slide_count = len(Presentation(filepath).slides)
                    return "num_of_slides", slide_count, False
                except Exception:
                    pass
            line_count = len(full_text.splitlines()) if full_text else 0
            return "num_of_text_lines", line_count, True

        if file_type == "pdf":
            if PdfReader is not None:
                try:
                    page_count = len(PdfReader(filepath).pages)
                    return "num_of_pages", page_count, False
                except Exception:
                    pass
            line_count = len(full_text.splitlines()) if full_text else 0
            estimated_pages = max(1, math.ceil(line_count / 45)) if line_count else 1
            return "num_of_pages", estimated_pages, True

        if file_type in {"docx", "doc"}:
            pages, _, estimated = self._count_docx_pages(filepath, full_text)
            return "num_of_pages", pages, estimated

        if file_type in text_like_extensions:
            return "num_of_text_lines", len(full_text.splitlines()) if full_text else 0, False

        return "num_of_text_lines", len(full_text.splitlines()) if full_text else 0, True

    def _process_file(self, filename, filepath):
        filepath = self._to_absolute_path(filepath)
        print(f"Szybkie skanowanie: {filename}...")
        stat = os.stat(filepath)
        file_type = filepath.split('.')[-1].lower()

        try:
            md_result = MarkItDown().convert(filepath)
            full_text = md_result.text_content
        except Exception as e:
            print(f"Błąd czytania {filename}: {e}")
            return None

        size_attr_name, size_attr_value, size_attr_estimated = self._infer_content_size(file_type, filepath, full_text)

        chunks = self.chunk_text(full_text)
        
        try:
            embed_response = ollama.embed(model=self.EMBED_MODEL, input=chunks)
            embeddings = [self._normalize_embedding(e) for e in embed_response['embeddings']]
        except Exception as e:
            print(f"Błąd wektoryzacji {filename}: {e}")
            return None

        document_record = {
            "filename": filename,
            "filepath": filepath,
            "metadata": {
                "file_type": file_type,
                "size_kb": round(stat.st_size / 1024, 2),
                size_attr_name: size_attr_value,
                "size_attr_is_estimated": size_attr_estimated,
                "created_at": datetime.fromtimestamp(stat.st_ctime).isoformat(),
                "status": "wektoryzacja_zakonczona" 
            },
            "ai_analysis": {
                "status": "oczekuje_na_analize", 
                "topic": None,
                "summary": None,
                "document_type": None,
                "keywords": [],
                "entities": {}
            },
            "chunks": [
                {"chunk_index": i, "content": c, "embedding": e} 
                for i, (c, e) in enumerate(zip(chunks, embeddings))
            ]
        }
        return document_record

    def process_folder(self):
        if not os.path.exists(self.folder_path): 
            print(f"Folder {self.folder_path} nie istnieje.")
            return

        files = [f for f in os.listdir(self.folder_path) if os.path.isfile(os.path.join(self.folder_path, f))]
        to_process = []

        for filename in files:
            filepath = self._to_absolute_path(os.path.join(self.folder_path, filename))
            if self.db.collection.find_one({"filepath": filepath}):
                continue 
            to_process.append((filename, filepath))

        if not to_process:
            print("Brak nowych plików do przetworzenia.")
            self.db.close()
            return

        with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
            futures = {executor.submit(self._process_file, fn, fp): fn for fn, fp in to_process}
            for future in as_completed(futures):
                doc_record = future.result()
                if doc_record:
                    self.db.collection.insert_one(doc_record)
                    print(f"Zapisano do bazy (bez AI): {doc_record['filename']}")
                
        self.db.close()
        print("\nZakończono etap 1: Pliki są gotowe do wektorowego wyszukiwania.")

if __name__ == "__main__":
    scanner = FastIngestionScanner()
    scanner.process_folder()